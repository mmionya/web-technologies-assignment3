"""Build the assignment's PDF and editable Word report from browser screenshots."""

import argparse
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PillowImage
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
SCREENSHOTS = ROOT / "report" / "screenshots"
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--name", default="Miras Zhumazhan")
parser.add_argument("--group", default="IT-2512")
parser.add_argument("--repository", default="https://github.com/mmionya/web-technologies-assignment3")
args = parser.parse_args()

# Both file formats use the same report content and original browser captures.
pages = [
    [
        ("title", "Assignment 3"),
        ("subtitle", "Responsive Web Design · Media Queries + Bootstrap Grid"),
        ("p", f"Name: {args.name}    |    Group: {args.group}"),
        ("p", f"Source repository: {args.repository}"),
        ("h", "Project and setup"),
        ("p", "This project is a fan-made Jamie Paige listening guide. The page includes an introduction, three reasons to listen, an album collection and a questions section. It uses local artwork and links to official music sources."),
        ("p", "index.html connects Bootstrap 5.3.8 CSS from its CDN, followed by style.css. The viewport meta tag sets the layout width to the device width. The Bootstrap JavaScript bundle appears before the closing body tag and powers the menu and accordion."),
        ("h", "Task 1 · A section with custom media queries"),
        ("p", "The #values section uses only custom classes. Its .values-grid starts with one column. At min-width: 768px it changes to two columns; at min-width: 1200px it changes to three. The media queries also adjust the heading size and spacing. This is mobile-first because the narrow layout works before either media query applies."),
        ("image", "values-1280.png", "Figure 1. Three custom CSS columns at a 1280 px viewport.", 6.65, 2.45),
        ("p", "Technical references: https://getbootstrap.com/docs/5.3/getting-started/introduction/ and https://getbootstrap.com/docs/5.3/layout/grid/. Artwork credits and official album sources are recorded in assets/CREDITS.md; README.md includes the component documentation links."),
    ],
    [
        ("h", "Task 1 · Phone and tablet evidence"),
        ("p", "At 375 px, each reason occupies its own row. At 768 px, two reasons share the first row and the third continues below. CSS Grid handles the remaining space without fixed item widths or absolute positioning. The screenshots below show the same section at two viewport widths."),
        ("images", [
            ("values-375.png", "Figure 2. 375 px: one column.", 2.05, 7.1),
            ("values-768.png", "Figure 3. 768 px: two columns.", 4.2, 7.1),
        ]),
    ],
    [
        ("h", "Task 2 · Bootstrap grid"),
        ("p", "The album section uses .container, .row and responsive columns: .col-12 .col-md-6 .col-xl-4. Bootstrap's 12-column grid therefore shows one album at 375 px, two at 768 px and three at 1280 px. Bootstrap cards keep the artwork, titles and links together; spacing utilities add consistent gaps."),
        ("image", "releases-1280.png", "Figure 4. Three Bootstrap album cards at 1280 px.", 6.65, 3.35),
        ("p", "The introduction uses two Bootstrap columns on wider screens. Responsive order-* classes put the introduction text before the artwork on a phone, so the purpose of the page appears first. On a desktop, the columns sit beside each other."),
        ("image", "about-1280.png", "Figure 5. The introduction's two-column desktop layout.", 6.65, 3.0),
    ],
    [
        ("h", "Task 2 · Mobile order and interactive components"),
        ("p", "The .navbar-expand-lg navigation collapses below 992 px. Its menu button reveals links to the page sections. The introduction's slogan uses .d-none .d-md-inline, hiding below 768 px; the decorative FAQ flower uses .d-none .d-lg-block. The FAQ is a Bootstrap accordion: its buttons reveal answers through data-bs-* attributes, so the page does not need custom JavaScript."),
        ("p", "These phone captures show the introduction with text first, the expanded menu, and an open answer. Browser checks also verify all three required viewport widths, both grids, section links, interactive controls and the absence of horizontal scrolling."),
        ("images", [
            ("about-375.png", "Figure 6. Text comes first at 375 px.", 2.12, 7.0),
            ("navbar-open-375.png", "Figure 7. The expanded mobile navigation.", 2.12, 7.0),
            ("accordion-open-375.png", "Figure 8. An expanded FAQ answer.", 2.12, 7.0),
        ]),
    ],
    [
        ("h", "Task 3 · Comparison of the two approaches"),
        ("images", [
            ("values-1280.png", "Custom CSS at 1280 px (detail: Figure 1).", 3.18, 1.3),
            ("releases-1280.png", "Bootstrap at 1280 px (detail: Figure 4).", 3.18, 1.3),
        ]),
        ("table", [
            ["Aspect", "Custom CSS: reasons to listen", "Bootstrap: albums and introduction"],
            ["Responsive code", "The .values-grid rule defines the base layout; two min-width media queries change its column count.", "The album columns use col-12 col-md-6 col-xl-4. Bootstrap already defines the breakpoint rules."],
            ["Design control", "The values section's heading size, gap and section spacing are adjusted directly in style.css.", "The grid supplies column widths and gutters. Custom CSS supplies the pink palette, type and card appearance."],
            ["Changing the layout", "To change the desktop count, edit grid-template-columns in the 1200 px media query.", "To change the desktop count, edit each album's col-xl-* class; the CSS framework supplies the widths."],
            ["Built-in behavior", "CSS Grid handles the repeating values, but does not provide an interactive menu or FAQ.", "The Bootstrap bundle handles navbar collapse and accordion state through data attributes."],
        ]),
        ("p", "I would choose custom media queries for a small section that needs its own spacing and layout rules, such as the reasons-to-listen grid. I would choose Bootstrap when several page sections need consistent responsive columns and standard components, as with the albums, navigation and FAQ. Combining them works here because the custom section stays independent while the rest reuses Bootstrap."),
        ("h", "Conclusion"),
        ("p", "This page shows how a mobile-first layout grows from one column to several columns as more space becomes available. The main difference is where the responsive rules live: in my own media queries or in Bootstrap's predefined classes. Testing at 375, 768 and 1280 px also makes it easier to catch ordering problems, overflow and controls that are difficult to use on a phone."),
        ("p", "Validation: tests/check_page.py passed at 320, 375, 768, 1024, 1280 and 1440 px, including both grids, section links, menu and accordion keyboard controls, image loading and no horizontal overflow."),
    ],
]
for width in (375, 768, 1280):
    pages.append([
        ("h", f"Whole-page screenshot · {width} px"),
        ("p", "Captured from the implemented page in a real browser. The entire page is included; use the task figures for close-up details."),
        ("image", f"page-{width}.png", f"Figure {9 + (375, 768, 1280).index(width)}. Complete page at a {width} px viewport.", 6.65, 8.45),
    ])


def image_size(filename, max_width, max_height):
    """Fit without stretching; all dimensions here are in inches."""
    with PillowImage.open(SCREENSHOTS / filename) as source:
        width, height = source.size
    scale = min(max_width / width, max_height / height)
    return width * scale, height * scale


# A missing capture fails clearly instead of creating a report with invented evidence.
for page in pages:
    for item in page:
        if item[0] == "image":
            assert (SCREENSHOTS / item[1]).is_file(), f"Missing screenshot: {item[1]}"
        elif item[0] == "images":
            for filename, *_ in item[1]:
                assert (SCREENSHOTS / filename).is_file(), f"Missing screenshot: {filename}"

stem = "Assignment3_" + "_".join(args.name.split())
output = ROOT / "report" / stem
document = Document()
document.core_properties.author = args.name
document.core_properties.title = "Assignment 3 — Responsive Web Design"
section = document.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.6)
section.left_margin = section.right_margin = Inches(0.7)
normal = document.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(10)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for heading in ("Title", "Heading 1", "Subtitle"):
    document.styles[heading].font.name = "Calibri"
    document.styles[heading].font.color.rgb = RGBColor.from_string("35252B")
document.styles["Title"].font.size = Pt(25)
document.styles["Heading 1"].font.size = Pt(16)
document.styles["Heading 1"].paragraph_format.space_before = Pt(8)
document.styles["Heading 1"].paragraph_format.space_after = Pt(8)
document.styles["Subtitle"].font.size = Pt(12)
document.styles["Caption"].font.size = Pt(8)
document.styles["Caption"].font.color.rgb = RGBColor.from_string("655B60")
document.styles["Caption"].paragraph_format.space_after = Pt(6)
section.footer.paragraphs[0].text = f"Assignment 3 | {args.name} | {args.group}"
section.footer.paragraphs[0].style = "Caption"

styles = getSampleStyleSheet()
styles.add(ParagraphStyle("Body", fontName="Helvetica", fontSize=10, leading=14, spaceAfter=9))
styles.add(ParagraphStyle("CaptionSmall", fontSize=8, leading=11, textColor=colors.HexColor("#655b60"), spaceAfter=9))
styles.add(ParagraphStyle("TableBody", fontSize=8.3, leading=11))
styles["Title"].fontSize, styles["Title"].leading = 26, 31
styles["Heading1"].fontSize, styles["Heading1"].leading = 16, 20
styles["Heading1"].spaceBefore, styles["Heading1"].spaceAfter = 8, 10
pdf = []

for page_number, page in enumerate(pages):
    if page_number:
        document.add_page_break()
        pdf.append(PageBreak())
    for item in page:
        kind = item[0]
        if kind in ("title", "subtitle", "h", "p"):
            docx_style = {"title": "Title", "subtitle": "Subtitle", "h": "Heading 1", "p": "Normal"}[kind]
            pdf_style = {"title": "Title", "subtitle": "Heading2", "h": "Heading1", "p": "Body"}[kind]
            document.add_paragraph(item[1], docx_style)
            pdf.append(Paragraph(escape(item[1]), styles[pdf_style]))
        elif kind == "image":
            _, filename, caption, max_width, max_height = item
            width, height = image_size(filename, max_width, max_height)
            document.add_picture(str(SCREENSHOTS / filename), width=Inches(width), height=Inches(height))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_paragraph(caption, "Caption")
            pdf.extend([Image(str(SCREENSHOTS / filename), width=width * inch, height=height * inch), Paragraph(escape(caption), styles["CaptionSmall"])])
        elif kind == "images":
            table = document.add_table(rows=1, cols=len(item[1]))
            table.autofit = False
            cells = []
            column_widths = []
            for column, cell, (filename, caption, max_width, max_height) in zip(table.columns, table.rows[0].cells, item[1]):
                width, height = image_size(filename, max_width, max_height)
                column.width = Inches(max_width + 0.12)
                cell.width = Inches(max_width + 0.12)
                cell.paragraphs[0].add_run().add_picture(str(SCREENSHOTS / filename), width=Inches(width), height=Inches(height))
                cell.add_paragraph(caption, "Caption")
                cells.append([Image(str(SCREENSHOTS / filename), width=width * inch, height=height * inch), Paragraph(escape(caption), styles["CaptionSmall"])])
                column_widths.append((max_width + 0.12) * inch)
            pdf_table = Table([cells], colWidths=column_widths)
            pdf_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 8)]))
            pdf.append(pdf_table)
        elif kind == "table":
            table = document.add_table(rows=0, cols=3)
            table.style = "Table Grid"
            table.autofit = False
            for column, width in zip(table.columns, (1.13, 2.76, 2.76)):
                column.width = Inches(width)
            for values in item[1]:
                for cell, value, width in zip(table.add_row().cells, values, (1.13, 2.76, 2.76)):
                    cell.width = Inches(width)
                    cell.text = value
                    cell.paragraphs[0].paragraph_format.space_after = Pt(6)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
            pdf_table = Table([[Paragraph(escape(value), styles["TableBody"]) for value in row] for row in item[1]], colWidths=[1.13 * inch, 2.76 * inch, 2.76 * inch])
            pdf_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f5e2e9")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d7c6cd")), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
            pdf.extend([pdf_table, Spacer(1, 12)])

document.save(str(output.with_suffix(".docx")))


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#655b60"))
    canvas.drawString(0.7 * inch, 0.32 * inch, f"Assignment 3 | {args.name} | {args.group}")
    canvas.drawRightString(7.57 * inch, 0.32 * inch, str(doc.page))
    canvas.restoreState()


SimpleDocTemplate(str(output.with_suffix(".pdf")), pagesize=(8.27 * inch, 11.69 * inch), rightMargin=0.7 * inch, leftMargin=0.7 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch, title="Assignment 3 — Responsive Web Design", author=args.name).build(pdf, onFirstPage=footer, onLaterPages=footer)
for suffix in (".pdf", ".docx"):
    assert output.with_suffix(suffix).stat().st_size > 0
    print(output.with_suffix(suffix))
